"""Smoke tests runnable without any API key:

    python tests/test_smoke.py

Covers: keyword SDG fallback, APA/cite-key helpers, markdown assembly,
word counting, and the .docx exporter.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config import PaperMeta
from src.docx_export import markdown_to_docx
from src.research import Source, _assign_cite_keys, build_digest, format_apa, publisher_for
from src.sdg_classifier import keyword_fallback
from src.writer import Section, assemble_markdown, word_count


def test_fallback_classifier():
    m = keyword_fallback("reducing plastic waste and recycling in Indian cities")
    assert m.method == "keyword-fallback"
    assert m.primary, "expected at least one primary SDG"
    # 11 (cities) and 12 (consumption/waste) should rank near the top
    assert any(n in {11, 12} for n in m.primary + m.secondary), m
    print("✔ fallback classifier:", m.primary, m.secondary)

    m2 = keyword_fallback("vaccination and maternal healthcare in rural districts")
    assert 3 in m2.primary or 3 in m2.secondary
    print("✔ healthcare topic mapped to SDG 3:", m2.primary)


def test_sources_and_digest():
    s1 = Source(
        title="Turning the tide on plastic: policy review",
        url="https://www.unep.org/resources/report/plastic-policy",
        domain="unep.org",
        publisher=publisher_for("unep.org"),
        snippet="Marine plastic pollution has doubled...",
        text="x" * 2000,
    )
    s2 = Source(
        title="Municipal solid waste in India",
        url="https://niti.gov.in/sites/default/files/waste-report-2024.pdf",
        domain="niti.gov.in",
        publisher=publisher_for("niti.gov.in"),
        snippet="India generates roughly 62 million tonnes of waste annually.",
    )
    _assign_cite_keys([s1, s2])
    assert s1.cite_key.startswith("(UN Environment Programme"), s1.cite_key
    assert s2.cite_key.startswith("(NITI Aayog"), s2.cite_key
    assert "unep.org" in format_apa(s1)
    digest = build_digest([s1, s2])
    assert "[S1]" in digest and "[S2]" in digest and "Citation key" in digest
    print("✔ cite keys:", s1.cite_key, "|", s2.cite_key)


def test_assembly_and_docx(tmp: Path):
    meta = PaperMeta(
        topic="plastic waste management in Indian cities",
        title="From Waste to Worth: Rethinking Plastic Governance in Indian Cities",
        policy_area="Urban Solid Waste Management",
        country_region="India",
        committee_event="Youth Policy Conclave 2026",
        participant_name="Test Delegate",
        institution="Test University",
        keywords=["plastic waste", "circular economy"],
    )
    exec_s = Section(
        key="executive_summary", number="1", title="Executive Summary",
        markdown="## 1. Executive Summary\n\nIndia faces a mounting plastic waste crisis. This paper proposes a reform package."
    )
    body = Section(
        key="problem_analysis", number="2", title="Problem Analysis",
        markdown=(
            "## 2. Problem Analysis\n\n"
            "### Current Situation\n\n"
            "India generates **62 million tonnes** of waste annually (NITI Aayog, 2024). *Cities* bear the brunt.\n\n"
            "### Root Cause Analysis\n\n"
            "Using PESTLE analysis:\n\n"
            "- Weak enforcement of extended producer responsibility\n"
            "- Fragmented municipal financing\n\n"
            "**Indicative Budget Summary**\n\n"
            "| Component | Year 1 | Year 2 | Total |\n"
            "| --- | --- | --- | --- |\n"
            "| Collection infrastructure | 120 crore | 150 crore | 270 crore |\n"
            "| Digital monitoring | 15 crore | 20 crore | 35 crore |\n"
        ),
    )
    refs = Section(
        key="references", number="", title="References",
        markdown=(
            "## References\n\n"
            "NITI Aayog. (2024). *Municipal solid waste in India*. NITI Aayog. https://niti.gov.in/example\n"
        ),
    )
    md, wc = assemble_markdown(meta, exec_s, [body], None, refs)
    assert wc > 10
    assert "approximately" in md
    assert word_count(body.markdown) > 20

    out = tmp / "test-paper.docx"
    markdown_to_docx(meta, md, str(out))
    assert out.exists() and out.stat().st_size > 5000

    # round-trip check: the docx should contain our content
    from docx import Document

    doc = Document(str(out))
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "From Waste to Worth" in texts
    assert "Problem Analysis" in texts
    assert "Unicamp" not in texts  # sanity
    assert len(doc.tables) == 1, "expected one table"
    assert doc.tables[0].rows[0].cells[0].text == "Component"
    print(f"✔ docx export ok ({out.stat().st_size:,} bytes, {wc} body words)")


if __name__ == "__main__":
    import tempfile

    test_fallback_classifier()
    test_sources_and_digest()
    with tempfile.TemporaryDirectory() as d:
        test_assembly_and_docx(Path(d))
    print("\nALL SMOKE TESTS PASSED ✅")
