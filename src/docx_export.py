"""Export the assembled paper to a submission-ready .docx.

Renders the cover page from metadata, then walks a constrained markdown
subset (headings, paragraphs, bold/italic, bullet/numbered lists, pipe
tables) and lays it out with academic formatting: Times New Roman 12,
1.5 line spacing, justified text, hanging-indent references, numbered
pages in the footer.
"""

from __future__ import annotations

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .config import PaperMeta

FONT = "Times New Roman"


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _add_page_number(doc: Document) -> None:
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    run.font.name = FONT
    run.font.size = Pt(10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)     # noqa: SLF001
    run._r.append(instr)         # noqa: SLF001
    run._r.append(fld_end)       # noqa: SLF001


def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    sizes = {"Heading 1": 16, "Heading 2": 14, "Heading 3": 12, "Title": 22}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = size >= 12
        style.font.color.rgb = RGBColor(0, 0, 0)


def _add_runs(paragraph, text: str) -> None:
    """Inline markdown: **bold**, *italic*, `code`, [text](url)."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)  # links -> text
    token_re = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    for chunk in token_re.split(text):
        if not chunk:
            continue
        run = paragraph.add_run()
        run.font.name = FONT
        if chunk.startswith("**") and chunk.endswith("**"):
            run.bold = True
            run.text = chunk[2:-2]
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run.italic = True
            run.text = chunk[1:-1]
        elif chunk.startswith("`") and chunk.endswith("`"):
            run.text = chunk[1:-1]
        else:
            run.text = chunk


def _clean_cell(text: str) -> str:
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.strip()


def _is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\|?[\s:|-]+\|[\s:|-]*$", line)) and "-" in line


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = doc.styles["Table Grid"]
    table.autofit = True
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncols):
            value = _clean_cell(row[j]) if j < len(row) else ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.name = FONT
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear")
                shade.set(qn("w:fill"), "E7E6E6")
                cells[j]._tc.get_or_add_tcPr().append(shade)  # noqa: SLF001
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Cover page
# ---------------------------------------------------------------------------

def _add_cover(
    doc: Document,
    meta: PaperMeta,
    subtitle: str = "A Policy Proposal",
    cover_fields: list[tuple[str, str]] | None = None,
) -> None:
    for _ in range(4):
        doc.add_paragraph()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(meta.title_or_default())
    run.font.name = FONT
    run.font.size = Pt(22)
    run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(subtitle)
    run.font.name = FONT
    run.font.size = Pt(13)
    run.italic = True

    doc.add_paragraph()
    if cover_fields is None:
        cover_fields = [
            ("Policy Area", meta.policy_area or "—"),
            ("Country/Region", meta.country_region or "—"),
            ("Committee/Event", meta.committee_event or "—"),
            ("Participant Name", meta.participant_name or "—"),
            ("Institution", meta.institution or "—"),
            ("Word Count", f"approximately {meta.word_count:,} words"),
            ("Keywords", ", ".join(meta.keywords) if meta.keywords else meta.topic),
        ]
    for label, value in cover_fields:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        label_run = p.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.name = FONT
        label_run.font.size = Pt(12)
        value_run = p.add_run(str(value))
        value_run.font.name = FONT
        value_run.font.size = Pt(12)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Main converter
# ---------------------------------------------------------------------------

def markdown_to_docx(
    meta: PaperMeta,
    markdown: str,
    out_path: str,
    cover_fields: list[tuple[str, str]] | None = None,
    subtitle: str = "A Policy Proposal",
) -> str:
    doc = Document()
    _setup_styles(doc)
    _add_page_number(doc)
    _add_cover(doc, meta, subtitle=subtitle, cover_fields=cover_fields)

    lines = markdown.splitlines()
    i = 0
    in_references = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if set(stripped) <= {"-", "=", "*", "_"} and len(stripped) >= 3:
            i += 1
            continue  # horizontal rule -> skip (cover already separated)

        # tables -------------------------------------------------------------
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for bl in block:
                if _is_table_sep(bl):
                    continue
                cells = [c.strip() for c in bl.strip("|").split("|")]
                rows.append(cells)
            _add_table(doc, rows)
            continue

        # headings -----------------------------------------------------------
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level <= 2:
                in_references = text.lower().startswith("references")
                doc.add_heading(text, level=1)
            elif level == 3:
                doc.add_heading(text, level=2)
            else:
                doc.add_heading(text, level=3)
            i += 1
            continue

        # lists ---------------------------------------------------------------
        m = re.match(r"^(\s*)[-•*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            _add_runs(p, m.group(2).strip())
            i += 1
            continue
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(2).strip())
            i += 1
            continue

        # paragraphs ----------------------------------------------------------
        p = doc.add_paragraph()
        if in_references:
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
        _add_runs(p, stripped)
        i += 1

    doc.save(out_path)
    return out_path
